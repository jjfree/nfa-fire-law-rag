"""Build synchronized DOCX and PDF editions of the frontend user guide."""

from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "FRONTEND_USER_GUIDE.md"
DOCX_OUTPUT = ROOT / "output" / "docx" / "nfa-fire-law-rag-frontend-user-guide.docx"
PDF_OUTPUT = ROOT / "output" / "pdf" / "nfa-fire-law-rag-frontend-user-guide.pdf"


@dataclass
class Block:
    kind: str
    text: str = ""
    level: int = 0
    rows: list[list[str]] | None = None
    path: Path | None = None
    caption: str = ""


def parse_source(path: Path) -> tuple[dict[str, str], list[Block]]:
    lines = path.read_text(encoding="utf-8").splitlines()
    if not lines or lines[0] != "---":
        raise ValueError("Guide source must start with YAML-like front matter")
    metadata: dict[str, str] = {}
    index = 1
    while index < len(lines) and lines[index] != "---":
        key, value = lines[index].split(":", 1)
        metadata[key.strip()] = value.strip()
        index += 1
    index += 1
    blocks: list[Block] = []
    paragraph: list[str] = []

    def flush_paragraph() -> None:
        if paragraph:
            blocks.append(
                Block("paragraph", text=" ".join(value.strip() for value in paragraph))
            )
            paragraph.clear()

    while index < len(lines):
        line = lines[index]
        if not line.strip():
            flush_paragraph()
            index += 1
            continue
        heading = re.match(r"^(#{2,3})\s+(.+)$", line)
        if heading:
            flush_paragraph()
            blocks.append(
                Block("heading", text=heading.group(2), level=len(heading.group(1)))
            )
            index += 1
            continue
        image = re.match(r"^!\[([^]]+)]\(([^)]+)\)$", line)
        if image:
            flush_paragraph()
            blocks.append(
                Block(
                    "image",
                    path=(path.parent / image.group(2)).resolve(),
                    caption=image.group(1),
                )
            )
            index += 1
            continue
        if line.startswith("```"):
            flush_paragraph()
            index += 1
            code: list[str] = []
            while index < len(lines) and not lines[index].startswith("```"):
                code.append(lines[index])
                index += 1
            blocks.append(Block("code", text="\n".join(code)))
            index += 1
            continue
        if (
            line.startswith("|")
            and index + 1 < len(lines)
            and re.match(r"^\|[-:|]+\|$", lines[index + 1].replace(" ", ""))
        ):
            flush_paragraph()
            table_lines = [line]
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            rows = [
                [cell.strip() for cell in row.strip().strip("|").split("|")]
                for row in table_lines
            ]
            blocks.append(Block("table", rows=rows))
            continue
        list_match = re.match(r"^([-*]|\d+\.)\s+(.+)$", line)
        if list_match:
            flush_paragraph()
            marker = list_match.group(1)
            kind = "numbered" if marker[0].isdigit() else "bullet"
            items: list[str] = []
            while index < len(lines):
                match = re.match(r"^([-*]|\d+\.)\s+(.+)$", lines[index])
                if not match:
                    break
                current_kind = "numbered" if match.group(1)[0].isdigit() else "bullet"
                if current_kind != kind:
                    break
                items.append(match.group(2))
                index += 1
            blocks.append(Block(kind, text="\n".join(items)))
            continue
        paragraph.append(line)
        index += 1
    flush_paragraph()
    return metadata, blocks


def set_run_font(
    run, name: str, size: float | None = None, bold: bool | None = None
) -> None:
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_pr.rFonts.set(qn("w:ascii"), name)
    r_pr.rFonts.set(qn("w:hAnsi"), name)
    r_pr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_markdown_runs(
    paragraph, text: str, font: str = "Microsoft JhengHei", size: float = 10.5
) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        code = part.startswith("`") and part.endswith("`")
        value = part[2:-2] if bold else part[1:-1] if code else part
        run = paragraph.add_run(value)
        set_run_font(run, "Consolas" if code else font, size=size, bold=bold)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    prefix = paragraph.add_run("第 ")
    set_run_font(prefix, "Microsoft JhengHei", 8)
    field_run = paragraph.add_run()
    set_run_font(field_run, "Microsoft JhengHei", 8)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run._r.extend([begin, instruction, end])
    suffix = paragraph.add_run(" 頁")
    set_run_font(suffix, "Microsoft JhengHei", 8)


def build_docx(metadata: dict[str, str], blocks: list[Block], output: Path) -> None:
    document = Document()
    document.core_properties.title = metadata["title"]
    document.core_properties.author = "台灣消防法規 RAG"
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft JhengHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.28

    title_style = document.styles["Title"]
    title_style.font.name = "Microsoft JhengHei"
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    title_style.font.size = Pt(24)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)

    for style_name, size in (("Heading 1", 16), ("Heading 2", 12.5)):
        style = document.styles[style_name]
        style.font.name = "Microsoft JhengHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    title = document.add_paragraph(style="Title")
    add_markdown_runs(title, metadata["title"], size=24)
    subtitle = document.add_paragraph()
    add_markdown_runs(subtitle, metadata["subtitle"], size=13)
    subtitle.paragraph_format.space_before = Pt(8)
    subtitle.paragraph_format.space_after = Pt(18)
    version = document.add_paragraph()
    add_markdown_runs(
        version,
        f"文件版本 {metadata['version']}\n更新日期 {metadata['date']}",
        size=10.5,
    )
    version.paragraph_format.line_spacing = 1.4
    lead = document.add_paragraph()
    add_markdown_runs(
        lead,
        "本手冊說明本機 Streamlit 前端的安裝啟動、日常查詢、對話管理、證據查核與常見問題。",
        size=11,
    )
    lead.paragraph_format.space_before = Pt(24)
    document.add_page_break()

    for block in blocks:
        if block.kind == "heading":
            style_name = "Heading 1" if block.level == 2 else "Heading 2"
            paragraph = document.add_paragraph(style=style_name)
            add_markdown_runs(
                paragraph, block.text, size=16 if block.level == 2 else 12.5
            )
        elif block.kind == "paragraph":
            paragraph = document.add_paragraph()
            add_markdown_runs(paragraph, block.text)
            if block.text.startswith("圖 "):
                paragraph.style = document.styles["Caption"]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(9)
                for run in paragraph.runs:
                    set_run_font(run, "Microsoft JhengHei", 9)
        elif block.kind in {"bullet", "numbered"}:
            style_name = "List Bullet" if block.kind == "bullet" else "List Number"
            for item in block.text.splitlines():
                paragraph = document.add_paragraph(style=style_name)
                add_markdown_runs(paragraph, item)
                paragraph.paragraph_format.space_after = Pt(3)
        elif block.kind == "code":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.15)
            paragraph.paragraph_format.right_indent = Inches(0.15)
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(8)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "F3F4F6")
            paragraph._p.get_or_add_pPr().append(shading)
            run = paragraph.add_run(block.text)
            set_run_font(run, "Consolas", 8.5)
        elif block.kind == "image" and block.path:
            if not block.path.exists():
                raise FileNotFoundError(block.path)
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            width = Inches(2.9) if "sidebar" in block.path.name else Inches(6.65)
            shape = paragraph.add_run().add_picture(str(block.path), width=width)
            shape._inline.docPr.set("descr", block.caption)
            paragraph.paragraph_format.keep_with_next = True
        elif block.kind == "table" and block.rows:
            rows = block.rows
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            for row_index, row in enumerate(rows):
                for col_index, value in enumerate(row):
                    cell = table.cell(row_index, col_index)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    set_cell_borders(cell)
                    if row_index == 0:
                        set_cell_shading(cell, "245B7A")
                    elif row_index % 2 == 0:
                        set_cell_shading(cell, "F3F6F8")
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(0)
                    add_markdown_runs(paragraph, value, size=8.6)
                    if row_index == 0:
                        for run in paragraph.runs:
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
            document.add_paragraph().paragraph_format.space_after = Pt(1)

    footer = section.footer.paragraphs[0]
    footer.text = ""
    add_page_number(footer)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(output)


def register_pdf_fonts() -> tuple[str, str]:
    candidates = [
        (Path(r"C:\Windows\Fonts\msjh.ttc"), Path(r"C:\Windows\Fonts\msjhbd.ttc")),
        (Path(r"C:\Windows\Fonts\mingliu.ttc"), Path(r"C:\Windows\Fonts\mingliub.ttc")),
    ]
    for regular, bold in candidates:
        if regular.exists():
            pdfmetrics.registerFont(
                TTFont("FrontendGuideRegular", str(regular), subfontIndex=0)
            )
            pdfmetrics.registerFont(
                TTFont(
                    "FrontendGuideBold",
                    str(bold if bold.exists() else regular),
                    subfontIndex=0,
                )
            )
            return "FrontendGuideRegular", "FrontendGuideBold"
    raise FileNotFoundError("No Traditional Chinese Windows font was found")


def pdf_markup(text: str, regular_font: str) -> str:
    value = escape(text, quote=False)
    value = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", value)
    value = re.sub(
        r"`([^`]+)`", rf'<font name="{regular_font}">\1</font>', value
    )
    return value


def build_pdf(metadata: dict[str, str], blocks: list[Block], output: Path) -> None:
    regular, bold = register_pdf_fonts()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "GuideTitle",
        parent=styles["Title"],
        fontName=bold,
        fontSize=24,
        leading=31,
        textColor=colors.black,
        alignment=TA_LEFT,
        spaceAfter=14,
    )
    subtitle_style = ParagraphStyle(
        "GuideSubtitle",
        parent=styles["BodyText"],
        fontName=regular,
        fontSize=13,
        leading=20,
        textColor=colors.black,
        spaceAfter=16,
    )
    h1 = ParagraphStyle(
        "GuideH1",
        parent=styles["Heading1"],
        fontName=bold,
        fontSize=15,
        leading=21,
        textColor=colors.black,
        spaceBefore=10,
        spaceAfter=6,
        keepWithNext=True,
    )
    h2 = ParagraphStyle(
        "GuideH2",
        parent=styles["Heading2"],
        fontName=bold,
        fontSize=11.5,
        leading=16,
        textColor=colors.black,
        spaceBefore=8,
        spaceAfter=4,
        keepWithNext=True,
    )
    body = ParagraphStyle(
        "GuideBody",
        parent=styles["BodyText"],
        fontName=regular,
        fontSize=9.6,
        leading=15,
        textColor=colors.HexColor("#202124"),
        spaceAfter=6,
    )
    bullet_style = ParagraphStyle(
        "GuideBullet", parent=body, fontSize=9.4, leading=14, spaceAfter=2
    )
    code_style = ParagraphStyle(
        "GuideCode",
        parent=styles["Code"],
        fontName=regular,
        fontSize=7.7,
        leading=10.5,
        leftIndent=7,
        rightIndent=7,
        textColor=colors.HexColor("#1F2937"),
        backColor=colors.HexColor("#F3F4F6"),
        borderPadding=6,
    )
    caption_style = ParagraphStyle(
        "GuideCaption",
        parent=body,
        fontSize=8.7,
        leading=12,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#4B5563"),
        spaceBefore=3,
        spaceAfter=8,
    )
    table_body = ParagraphStyle(
        "GuideTable", parent=body, fontSize=8, leading=11, spaceAfter=0
    )
    table_header = ParagraphStyle(
        "GuideTableHeader",
        parent=table_body,
        fontName=bold,
        textColor=colors.white,
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    document = SimpleDocTemplate(
        str(output),
        pagesize=LETTER,
        leftMargin=0.68 * inch,
        rightMargin=0.68 * inch,
        topMargin=0.64 * inch,
        bottomMargin=0.7 * inch,
        title=metadata["title"],
        author="台灣消防法規 RAG",
    )
    available_width = LETTER[0] - document.leftMargin - document.rightMargin
    story: list[object] = [
        Spacer(1, 0.55 * inch),
        Paragraph(pdf_markup(metadata["title"], regular), title_style),
        Paragraph(pdf_markup(metadata["subtitle"], regular), subtitle_style),
        Spacer(1, 0.16 * inch),
        Paragraph(
            f"文件版本 {metadata['version']}<br/>更新日期 {metadata['date']}", body
        ),
        Spacer(1, 0.35 * inch),
        Paragraph(
            "本手冊說明本機 Streamlit 前端的安裝啟動、日常查詢、對話管理、證據查核與常見問題。",
            subtitle_style,
        ),
        PageBreak(),
    ]
    for block in blocks:
        if block.kind == "heading":
            story.append(
                Paragraph(
                    pdf_markup(block.text, regular), h1 if block.level == 2 else h2
                )
            )
        elif block.kind == "paragraph":
            paragraph_style = caption_style if block.text.startswith("圖 ") else body
            story.append(
                Paragraph(pdf_markup(block.text, regular), paragraph_style)
            )
        elif block.kind in {"bullet", "numbered"}:
            items = [
                ListItem(Paragraph(pdf_markup(item, regular), bullet_style))
                for item in block.text.splitlines()
            ]
            story.append(
                ListFlowable(
                    items,
                    bulletType="1" if block.kind == "numbered" else "bullet",
                    leftIndent=18,
                    bulletFontName=regular,
                    bulletFontSize=8,
                )
            )
            story.append(Spacer(1, 3))
        elif block.kind == "code":
            story.extend(
                [Preformatted(block.text, code_style, maxLineLength=100), Spacer(1, 6)]
            )
        elif block.kind == "image" and block.path:
            if not block.path.exists():
                raise FileNotFoundError(block.path)
            with PILImage.open(block.path) as source_image:
                source_width, source_height = source_image.size
            max_width = 2.9 * inch if "sidebar" in block.path.name else available_width
            max_height = 6.6 * inch
            scale = min(max_width / source_width, max_height / source_height)
            picture = Image(
                str(block.path),
                width=source_width * scale,
                height=source_height * scale,
            )
            picture.hAlign = "CENTER"
            story.append(picture)
        elif block.kind == "table" and block.rows:
            data = [
                [
                    Paragraph(
                        pdf_markup(cell, regular),
                        table_header if row_index == 0 else table_body,
                    )
                    for cell in row
                ]
                for row_index, row in enumerate(block.rows)
            ]
            column_count = len(block.rows[0])
            proportions = (
                [0.20, 0.35, 0.45]
                if column_count == 3
                else [0.30, 0.70]
                if column_count == 2
                else [1 / column_count] * column_count
            )
            table = Table(
                data,
                colWidths=[available_width * value for value in proportions],
                repeatRows=1,
                hAlign="LEFT",
            )
            table.setStyle(
                TableStyle(
                    [
                        (
                            "BACKGROUND",
                            (0, 0),
                            (-1, 0),
                            colors.HexColor("#245B7A"),
                        ),
                        (
                            "ROWBACKGROUNDS",
                            (0, 1),
                            (-1, -1),
                            [colors.white, colors.HexColor("#F3F6F8")],
                        ),
                        (
                            "GRID",
                            (0, 0),
                            (-1, -1),
                            0.35,
                            colors.HexColor("#D9D9D9"),
                        ),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 6),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                        ("TOPPADDING", (0, 0), (-1, -1), 5),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ]
                )
            )
            story.extend([table, Spacer(1, 7)])

    def footer(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont(regular, 7.5)
        canvas.setFillColor(colors.HexColor("#6B7280"))
        canvas.drawString(document.leftMargin, 0.35 * inch, metadata["title"])
        canvas.drawRightString(
            LETTER[0] - document.rightMargin, 0.35 * inch, f"第 {doc.page} 頁"
        )
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    print(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--format", choices=("docx", "pdf", "all"), default="all"
    )
    args = parser.parse_args()
    metadata, blocks = parse_source(SOURCE)
    if args.format in {"docx", "all"}:
        build_docx(metadata, blocks, DOCX_OUTPUT)
    if args.format in {"pdf", "all"}:
        build_pdf(metadata, blocks, PDF_OUTPUT)


if __name__ == "__main__":
    main()
